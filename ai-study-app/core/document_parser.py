"""
Bộ phân tích tài liệu đa định dạng: PDF, DOCX, PPTX, IPYNB
"""
import io
import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def parse_pdf(file_bytes: bytes) -> str:
    """Trích xuất văn bản từ file PDF."""
    text_parts = []
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for i, page in enumerate(pdf.pages, 1):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(f"[Trang {i}]\n{page_text.strip()}")
        if text_parts:
            return "\n\n".join(text_parts)
    except Exception as e:
        logger.warning(f"pdfplumber thất bại: {e}, thử PyPDF2...")

    try:
        import PyPDF2
        reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        for i, page in enumerate(reader.pages, 1):
            page_text = page.extract_text()
            if page_text:
                text_parts.append(f"[Trang {i}]\n{page_text.strip()}")
        return "\n\n".join(text_parts)
    except Exception as e:
        logger.error(f"Không thể đọc PDF: {e}")
        return ""


def parse_docx(file_bytes: bytes) -> str:
    """Trích xuất văn bản từ file DOCX."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        parts = []

        for para in doc.paragraphs:
            if para.text.strip():
                style = para.style.name if para.style else ""
                prefix = "## " if "Heading 1" in style else "### " if "Heading 2" in style else ""
                parts.append(f"{prefix}{para.text.strip()}")

        # Extract tables
        for table in doc.tables:
            table_text = []
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    table_text.append(row_text)
            if table_text:
                parts.append("[Bảng]\n" + "\n".join(table_text))

        return "\n\n".join(parts)
    except Exception as e:
        logger.error(f"Không thể đọc DOCX: {e}")
        return ""


def parse_pptx(file_bytes: bytes) -> str:
    """Trích xuất văn bản từ file PPTX."""
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(file_bytes))
        parts = []

        for i, slide in enumerate(prs.slides, 1):
            slide_texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_texts.append(shape.text.strip())
            if slide_texts:
                parts.append(f"[Slide {i}]\n" + "\n".join(slide_texts))

        return "\n\n".join(parts)
    except Exception as e:
        logger.error(f"Không thể đọc PPTX: {e}")
        return ""


def parse_ipynb(file_bytes: bytes) -> str:
    """Trích xuất markdown + code logic từ Jupyter Notebook."""
    try:
        import nbformat
        nb = nbformat.reads(file_bytes.decode("utf-8", errors="replace"), as_version=4)
        parts = []

        for i, cell in enumerate(nb.cells, 1):
            source = cell.get("source", "").strip()
            if not source:
                continue
            if cell.cell_type == "markdown":
                parts.append(f"[Ghi chú Markdown - Cell {i}]\n{source}")
            elif cell.cell_type == "code":
                parts.append(f"[Code - Cell {i}]\n```python\n{source}\n```")
                # Include output if text-based
                outputs = cell.get("outputs", [])
                for out in outputs:
                    if out.get("output_type") in ("stream", "execute_result"):
                        out_text = "".join(out.get("text", out.get("data", {}).get("text/plain", [])))
                        if out_text.strip():
                            parts.append(f"[Output Cell {i}]\n{out_text.strip()[:500]}")

        return "\n\n".join(parts)
    except Exception as e:
        logger.error(f"Không thể đọc IPYNB: {e}")
        return ""


def parse_text(file_bytes: bytes) -> str:
    """Đọc file văn bản thuần (.txt, .md) với nhiều bảng mã."""
    for enc in ("utf-8-sig", "utf-16", "cp1258", "cp1252"):
        try:
            return file_bytes.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return file_bytes.decode("utf-8", errors="replace")


def parse_file(filename: str, file_bytes: bytes) -> str:
    """
    Dispatcher chính: chọn parser phù hợp theo đuôi file.

    Returns:
        Chuỗi văn bản đã trích xuất, hoặc chuỗi rỗng nếu thất bại.
        Không cắt bớt nội dung: việc chia nhỏ do llm_client đảm nhiệm.
    """
    ext = Path(filename).suffix.lower()
    parsers = {
        ".pdf":   parse_pdf,
        ".docx":  parse_docx,
        ".pptx":  parse_pptx,
        ".ipynb": parse_ipynb,
        ".txt":   parse_text,
        ".md":    parse_text,
    }

    parser = parsers.get(ext)
    if not parser:
        logger.warning(f"Định dạng không hỗ trợ: {ext}")
        return ""

    content = parser(file_bytes)
    if not content.strip():
        logger.warning(f"Không trích xuất được nội dung từ {filename}")
        return ""
    return content
